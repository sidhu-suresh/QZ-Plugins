#!/usr/bin/env python3
"""
deck_check.py - Quantzig deck sweep. Run on every finished deck, then fix and re-run until clean.

  python3 deck_check.py deck.pptx                     full report, slide by slide
  python3 deck_check.py deck.pptx --fix -o fixed.pptx fix fonts to Grandview and font sizes to allowed values

Checks
  FONT       any text font that is not Grandview (FontAwesome allowed for icon glyphs only)
  SIZE       any font size that is not a whole number of points, or below 10pt. 10.5pt is the one allowed half size
  COLOUR     off-palette colours (always a defect)
  STATUS     green / amber / red uses: allowed only on metric status, confirm each one
  WORDS      content slides with too little copy to be read without a voice-over (default under 70 words)
  SPACE      content slides where shapes cover too little of the content area (default under 70%)
  SAME       pairs of slides whose layouts are too alike, and library visuals used more than once
"""
import argparse, re, shutil, sys, zipfile
from collections import defaultdict

from pptx import Presentation
from pptx.util import Emu

BRAND = {"0B0E5F", "812B47", "7F2A4C", "2C165E", "592258", "242424", "000000", "FFFFFF",
         "F8F4F6", "EDD3D4", "D2D0E1"}
STATUS = {"30A050": "green", "C38424": "amber", "BD413A": "red",
          "F7FDF9": "mint tint", "FDFBF6": "cream tint", "FDF7F7": "blush tint"}
FONT_OK = {"grandview", "grandview display"}
KEEP_FONTS = ("awesome", "wingding", "symbol")
FONT_TAGS = ("latin", "ea", "cs", "buFont", "sym")
EMU_IN = 914400
NON_CONTENT_LAYOUTS = ("title", "cover", "partition", "agenda")


def allowed_size(sz):
    return sz >= 1000 and (sz % 100 == 0 or sz == 1050)


def fix_size(sz):
    if sz < 1000:
        return 1000
    if sz == 1050 or sz % 100 == 0:
        return sz
    if 1000 < sz < 1100:
        return 1050 if abs(sz - 1050) <= 25 else (1000 if sz < 1050 else 1100)
    return int(round(sz / 100.0)) * 100


def xml_checks(xml, is_slide):
    out = defaultdict(dict)
    for tag in FONT_TAGS:
        for m in re.finditer(rf'<a:{tag}\b[^>]*typeface="([^"]*)"', xml):
            f = m.group(1)
            if not f or f.startswith("+") or f.lower() in FONT_OK or any(k in f.lower() for k in KEEP_FONTS):
                continue
            out["FONT"][f] = out["FONT"].get(f, 0) + 1
    for m in re.finditer(r'<a:(?:rPr|defRPr|endParaRPr)\b[^>]*\bsz="(\d+)"', xml):
        sz = int(m.group(1))
        if not allowed_size(sz):
            k = f"{sz/100:g}pt"
            out["SIZE"][k] = out["SIZE"].get(k, 0) + 1
    if is_slide:
        for m in re.finditer(r'<a:srgbClr val="([0-9A-Fa-f]{6})"', xml):
            v = m.group(1).upper()
            if v in BRAND:
                continue
            d = "STATUS" if v in STATUS else "COLOUR"
            key = f"{STATUS.get(v, '#' + v)}" if d == "STATUS" else "#" + v
            out[d][key] = out[d].get(key, 0) + 1
        body = re.sub(r"<p:style>.*?</p:style>", "", xml, flags=re.S)
        for m in re.finditer(r'<a:schemeClr val="(accent5|accent6)"', body):
            key = "amber (accent5)" if m.group(1) == "accent5" else "green (accent6)"
            out["STATUS"][key] = out["STATUS"].get(key, 0) + 1
    return out


def fix_xml(xml):
    def font(m):
        f = m.group(2)
        if f.startswith("+") or f.lower() in FONT_OK or any(k in f.lower() for k in KEEP_FONTS):
            return m.group(0)
        return f'{m.group(1)}typeface="Grandview"'
    xml = re.sub(r'(<a:(?:%s)\b[^>]*?)typeface="([^"]*)"' % "|".join(FONT_TAGS), font, xml)

    def size(m):
        return f'{m.group(1)}sz="{fix_size(int(m.group(2)))}"'
    return re.sub(r'(<a:(?:rPr|defRPr|endParaRPr)\b[^>]*?\b)sz="(\d+)"', size, xml)


def is_content(slide, idx, n):
    name = slide.slide_layout.name.lower()
    if idx in (0, n - 1):
        return False
    return not any(k in name for k in NON_CONTENT_LAYOUTS)


def slide_words(slide):
    words = 0
    for sh in slide.shapes:
        words += shape_words(sh)
    return words


def shape_words(sh):
    if sh.is_placeholder and sh.placeholder_format.type in (13, 15, 16):  # slide number, footer, date
        return 0
    n = 0
    if sh.shape_type == 6:  # group
        for s in sh.shapes:
            n += shape_words(s)
    if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
        n += len(re.findall(r"[A-Za-z0-9%$]+", sh.text_frame.text))
    if getattr(sh, "has_table", False) and sh.has_table:
        for row in sh.table.rows:
            for c in row.cells:
                n += len(re.findall(r"[A-Za-z0-9%$]+", c.text))
    return n


def coverage(slide, W, H):
    # content area: below the header ribbon, above the footer, inside side margins
    x0, x1 = 0.35 * EMU_IN, W - 0.35 * EMU_IN
    y0, y1 = 1.3 * EMU_IN, H - 0.75 * EMU_IN
    step = int(0.1 * EMU_IN)
    cols, rows = int((x1 - x0) / step), int((y1 - y0) / step)
    grid = [[False] * cols for _ in range(rows)]
    for sh in slide.shapes:
        if sh.is_placeholder and sh.placeholder_format.type in (1, 3, 13, 15, 16):
            continue
        if sh.left is None or sh.width is None:
            continue
        for r in range(rows):
            cy = y0 + r * step + step / 2
            if not (sh.top <= cy <= sh.top + sh.height):
                continue
            for c in range(cols):
                cx = x0 + c * step + step / 2
                if sh.left <= cx <= sh.left + sh.width:
                    grid[r][c] = True
    tot = rows * cols
    return 100.0 * sum(map(sum, grid)) / tot if tot else 100.0


def signature(slide):
    sig = set()
    for sh in slide.shapes:
        if sh.is_placeholder or sh.left is None or sh.width is None:
            continue
        kind = sh.shape_type
        try:
            geom = sh.auto_shape_type if kind == 1 else None
        except Exception:
            geom = None
        q = lambda v: int(round(v / (0.6 * EMU_IN)))
        sig.add((str(kind), str(geom), q(sh.left), q(sh.top), q(sh.width), q(sh.height)))
    return sig


def library_ids(slide):
    ids = []
    for sh in slide.shapes:
        m = re.match(r"Visual L(\d+)", sh.name or "")
        if m:
            ids.append(int(m.group(1)))
    return ids


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("--fix", action="store_true", help="fix fonts and font sizes in slides, layouts and master")
    ap.add_argument("-o", "--out")
    ap.add_argument("--min-words", type=int, default=70)
    ap.add_argument("--min-coverage", type=float, default=70.0)
    ap.add_argument("--same", type=float, default=0.55, help="layout similarity threshold (0-1)")
    a = ap.parse_args()

    if a.fix:
        out = a.out or a.pptx
        tmp = out + ".tmp"
        with zipfile.ZipFile(a.pptx) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if re.match(r"ppt/(slides|slideLayouts|slideMasters)/[^/]+\.xml$", item.filename):
                    data = fix_xml(data.decode("utf8")).encode("utf8")
                zout.writestr(item, data)
        shutil.move(tmp, out)
        print(f"fonts set to Grandview and sizes set to whole points (10.5 kept) -> {out}\n")
        a.pptx = out

    prs = Presentation(a.pptx)
    W, H = prs.slide_width, prs.slide_height
    z = zipfile.ZipFile(a.pptx)
    slides = list(prs.slides)
    n = len(slides)
    problems = 0
    sigs, libs = {}, defaultdict(list)
    for i, s in enumerate(slides):
        part = s.part.partname.lstrip("/")
        xml = z.read(part).decode("utf8")
        res = xml_checks(xml, True)
        lines = []
        for k in ("FONT", "SIZE", "COLOUR"):
            if res.get(k):
                problems += 1
                lines.append(f"  {k}: " + ", ".join(f"{a_} x{b}" for a_, b in res[k].items()))
        if res.get("STATUS"):
            lines.append("  STATUS (confirm each marks a metric): " + ", ".join(f"{a_} x{b}" for a_, b in res["STATUS"].items()))
        if is_content(s, i, n):
            w = slide_words(s)
            cov = coverage(s, W, H)
            if w < a.min_words:
                problems += 1
                lines.append(f"  WORDS: {w} words; too thin to read without a voice-over (target {a.min_words}+)")
            if cov < a.min_coverage:
                problems += 1
                lines.append(f"  SPACE: shapes cover {cov:.0f}% of the content area (target {a.min_coverage:.0f}%+)")
            sigs[i] = signature(s)
        for lid in library_ids(s):
            libs[lid].append(i + 1)
        if lines:
            print(f"Slide {i + 1} ({s.slide_layout.name}):")
            print("\n".join(lines))
    # layout sameness
    keys = sorted(sigs)
    same = []
    for x in range(len(keys)):
        for y in range(x + 1, len(keys)):
            A, B = sigs[keys[x]], sigs[keys[y]]
            if len(A) < 3 or len(B) < 3:
                continue
            j = len(A & B) / len(A | B)
            if j >= a.same:
                same.append((keys[x] + 1, keys[y] + 1, j))
    for lid, where in libs.items():
        if len(where) > 1:
            same.append((where[0], where[1], 1.0))
            print(f"SAME: library visual {lid} used on slides {', '.join(map(str, where))}; use a different visual on all but one")
    for s1, s2, j in same:
        if j < 1.0:
            print(f"SAME: slides {s1} and {s2} share {j:.0%} of their layout; rebuild one with a different structure")
    problems += len(same)
    # layouts and master: fonts and sizes only
    for part in sorted(p for p in z.namelist() if re.match(r"ppt/(slideLayouts|slideMasters)/[^/]+\.xml$", p)):
        res = xml_checks(z.read(part).decode("utf8"), False)
        if res.get("FONT") or res.get("SIZE"):
            problems += 1
            print(f"{part.replace('ppt/', '')}: " + "; ".join(
                f"{k}: " + ", ".join(f"{a_} x{b}" for a_, b in res[k].items()) for k in ("FONT", "SIZE") if res.get(k)))
    print(f"\n{'CLEAN' if not problems else f'{problems} issue(s) to fix'}. STATUS lines always need a human check.")


if __name__ == "__main__":
    main()
