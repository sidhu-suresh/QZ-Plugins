#!/usr/bin/env python3
"""
visuals.py - lift infographic/diagram visuals out of the bundled visual library
(assets/visual-library.pptx) and drop them into any deck, or render them as images.

Subcommands
  inspect  LIB_SLIDE                 list what is on a library slide (top shapes T#, texts t#, icons s#, image slots i#)
  preview  LIB_SLIDE -o out.png      render a library slide with t#/s#/i# labels drawn on it
  place    LIB_SLIDE --target deck.pptx (--slide N | --new-slide LAYOUT) [options] -o out.pptx
  render   LIB_SLIDE -o visual.png   [same editing options as place]  -> PNG for HTML / Word / PDF / artifacts

Editing options (place and render)
  --texts FILE.json      {"t3": "Discover", "t4": "**Scope:** agree KPIs"}  (or a JSON list in t-order; null = keep)
                         "\\n" starts a new paragraph; a leading **bold** segment becomes a bold run.
  --drop T2,T7,s14       remove top-level shapes (T#) or any nested shape (s#) you do not need
  --icons FILE.json      {"s12": "database", "s20": "line-chart"}  swap a library icon for a FontAwesome 4 icon
  --icon-mode png|glyph  png (default, renders anywhere) or glyph (FontAwesome text, recolourable, needs the font)
  --images a.png,b.png   fill image slots i0, i1, ... in order (e.g. dashboard screenshots into laptop mockups)
  --palette keep|quantzig|custom   colour handling (default keep = inherit target theme accents)
  --colors accent1=0B0E5F,accent2=812B47,...   hex per library accent slot (with --palette custom)
  --font NAME            replace the library's Roboto faces with NAME (defaults to Grandview with --palette quantzig)
  --min-font PT          raise any text below PT to PT after scaling (Quantzig: 10)
  --box X,Y,W,H          inches on the target slide to fit the visual into (aspect kept, centred)
"""
import argparse, copy, io, json, os, re, shutil, subprocess, sys, tempfile

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
LIB_DEFAULT = os.path.join(SKILL, "assets", "visual-library.pptx")
FA_TTF = os.path.join(SKILL, "assets", "fontawesome-webfont.ttf")
FA_MAP = os.path.join(SKILL, "references", "fontawesome4-icons.tsv")

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"p": P, "a": A, "r": R}
SHAPE_TAGS = {f"{{{P}}}{t}" for t in ("sp", "grpSp", "pic", "cxnSp", "graphicFrame")}
TXBODY_TAGS = {f"{{{P}}}txBody", f"{{{A}}}txBody"}
ACCENTS = [f"accent{i}" for i in range(1, 7)]
EMU_IN = 914400

# Library theme accents (for resolving colours when we need a hex, e.g. PNG icons)
LIB_ACCENTS = {"accent1": "237DB9", "accent2": "15AA96", "accent3": "9BB955",
               "accent4": "F19B14", "accent5": "BE382C", "accent6": "633247"}

QZ_NAVY, QZ_WINE, QZ_PURPLE = "0B0E5F", "812B47", "2C165E"
QZ_MAP = {"accent1": QZ_NAVY, "accent2": QZ_WINE, "accent3": QZ_PURPLE,
          "accent4": QZ_NAVY, "accent5": QZ_WINE, "accent6": QZ_PURPLE}

PLACEHOLDER_TEXT = re.compile(r"lorem|ipsum|title goes here|keyword|there are many variations|"
                              r"number (one|two|three|four|five|six)|goes here|just start|description", re.I)


def q(tag):
    pre, local = tag.split(":")
    return f"{{{NS[pre]}}}{local}"


# ----------------------------------------------------------------------------- geometry

def xfrm_of(el):
    if el.tag == q("p:grpSp"):
        return el.find("p:grpSpPr/a:xfrm", NS)
    if el.tag == q("p:graphicFrame"):
        return el.find("p:xfrm", NS)
    return el.find("p:spPr/a:xfrm", NS)


def rect_of(el):
    x = xfrm_of(el)
    if x is None:
        return None
    off, ext = x.find("a:off", NS), x.find("a:ext", NS)
    if off is None or ext is None:
        return None
    return (int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy")))


def walk(el, tf=(1.0, 1.0, 0.0, 0.0), out=None, depth=0):
    """Depth-first over shape elements, recording absolute (slide-space) rects."""
    if out is None:
        out = []
    sx, sy, ox, oy = tf
    r = rect_of(el)
    ab = None
    if r:
        ab = (ox + r[0] * sx, oy + r[1] * sy, r[2] * sx, r[3] * sy)
    out.append({"el": el, "abs": ab, "depth": depth})
    if el.tag == q("p:grpSp"):
        x = xfrm_of(el)
        choff, chext = x.find("a:chOff", NS), x.find("a:chExt", NS)
        if r and choff is not None and chext is not None and int(chext.get("cx")) and int(chext.get("cy")):
            ksx = r[2] / int(chext.get("cx"))
            ksy = r[3] / int(chext.get("cy"))
            ntf = (sx * ksx, sy * ksy,
                   ox + sx * (r[0] - int(choff.get("x")) * ksx),
                   oy + sy * (r[1] - int(choff.get("y")) * ksy))
        else:
            ntf = tf
        for ch in el:
            if ch.tag in SHAPE_TAGS:
                walk(ch, ntf, out, depth + 1)
    return out


def text_of(txbody):
    paras = []
    for p in txbody.findall("a:p", NS):
        s = ""
        for node in p:
            if node.tag in (q("a:r"), q("a:fld")):
                t = node.find("a:t", NS)
                s += t.text or "" if t is not None else ""
            elif node.tag == q("a:br"):
                s += " / "
        paras.append(s)
    return "\n".join(paras).strip()


def is_placeholder(el):
    return bool(etree._Element.xpath(el, "./*[1]/p:nvPr/p:ph", namespaces=NS))


def is_chart(el):
    return el.tag == q("p:graphicFrame") and bool(
        etree._Element.xpath(el, ".//a:graphicData[contains(@uri,'chart')]", namespaces=NS))


def name_of(el):
    c = el.find("./*[1]/p:cNvPr", NS)
    return c.get("name") if c is not None else "?"


# ----------------------------------------------------------------------------- library access

def lib_slide(prs, n):
    if n < 1 or n > len(prs.slides):
        sys.exit(f"library slide {n} out of range 1..{len(prs.slides)}")
    return prs.slides[n - 1]


def layout_ph_rect(slide, idx):
    for ph in slide.slide_layout.placeholders:
        if ph.placeholder_format.idx == idx:
            sp = ph._element.find("p:spPr", NS)
            geom = sp.find("a:prstGeom", NS) if sp is not None else None
            return (ph.left, ph.top, ph.width, ph.height), (geom.get("prst") if geom is not None else "rect")
    return None, "rect"


def collect(slide):
    """Top-level visual elements in z-order. Picture placeholders become image slots; other
    placeholders (title, subtitle, slide number) and charts are skipped."""
    tops, skipped = [], []
    # device frames, photo frames etc. live on some library layouts: bring them along (behind)
    for el in slide.slide_layout.shapes._spTree:
        if el.tag not in SHAPE_TAGS or is_placeholder(el) or is_chart(el):
            continue
        nm = name_of(el)
        if nm.startswith("Round Same Side Corner Rectangle") or nm.startswith("Slide Number"):
            continue  # the library's title tab / page number chrome
        tops.append({"el": el, "slot": False, "part": slide.slide_layout.part, "layout": True})
    for el in slide.shapes._spTree:
        if el.tag not in SHAPE_TAGS:
            continue
        if is_placeholder(el):
            ph = el.find("./*[1]/p:nvPr/p:ph", NS)
            if el.tag == q("p:pic") or ph.get("type") == "pic":
                idx = int(ph.get("idx", "0"))
                rect = rect_of(el)
                geom_el = el.find("p:spPr/a:prstGeom", NS)
                geom = geom_el.get("prst") if geom_el is not None else None
                if rect is None:
                    rect, lgeom = layout_ph_rect(slide, idx)
                    geom = geom or lgeom
                if rect:
                    tops.append({"el": el, "slot": True, "rect": rect, "geom": geom or "rect", "part": slide.part})
            continue
        if is_chart(el):
            skipped.append(f"chart '{name_of(el)}' (rebuild charts natively with real data)")
            continue
        tops.append({"el": el, "slot": False, "part": slide.part})
    return tops, skipped


def index(tops):
    """Stable labels: T# top-level, s# every shape (nested too), t# every non-empty text body,
    i# image slots, plus icon candidates."""
    shapes, texts, slots, icons = [], [], [], []
    for ti, top in enumerate(tops):
        top["T"] = ti
        if top["slot"]:
            slots.append({"i": len(slots), "top": top, "abs": top["rect"]})
            continue
        recs = walk(top["el"])
        for rec in recs:
            rec["s"] = len(shapes)
            rec["T"] = ti
            shapes.append(rec)
        for rec in recs:
            el = rec["el"]
            for tb in [c for c in el if c.tag in TXBODY_TAGS] + (
                    el.findall(".//a:tbl//a:txBody", NS) if el.tag == q("p:graphicFrame") else []):
                t = text_of(tb)
                if t:
                    texts.append({"t": len(texts), "tb": tb, "text": t, "abs": rec["abs"],
                                  "shape": name_of(el), "T": ti, "sz": first_size(tb)})
        # icon candidates: outermost text-free custom-geometry blobs of icon size
        taken = []
        for rec in recs:
            el, ab = rec["el"], rec["abs"]
            if ab is None or el.tag in (q("p:cxnSp"), q("p:graphicFrame"), q("p:pic")):
                continue
            if any(is_ancestor(t, el) for t in taken):
                continue
            w, h = ab[2] / EMU_IN, ab[3] / EMU_IN
            if not (0.1 <= w <= 0.95 and 0.1 <= h <= 0.95 and 0.4 <= (w / h if h else 0) <= 2.5):
                continue
            if any(text_of(tb) for tb in el.iter() if tb.tag in TXBODY_TAGS):
                continue
            has_cust = el.find(".//a:custGeom", NS) is not None
            if not has_cust:
                continue
            if el.tag == q("p:sp") and len(list(el.find(".//a:custGeom", NS).iter())) < 40:
                continue  # simple polygons (e.g. iceberg shards) are artwork, not icons
            # a coloured badge (circle/square) + glyph group: offer the glyph, keep the badge
            if el.tag == q("p:grpSp"):
                kids = [r2 for r2 in recs if r2["el"].getparent() is el]
                badge = None
                for k2 in kids:
                    g = k2["el"].find("p:spPr/a:prstGeom", NS)
                    if g is not None and k2["abs"] and k2["abs"][2] * k2["abs"][3] >= 0.6 * ab[2] * ab[3]:
                        badge = k2
                if badge is not None:
                    glyphs = [k2 for k2 in kids if k2 is not badge and k2["el"].find(".//a:custGeom", NS) is not None
                              or (k2 is not badge and k2["el"].tag == q("p:grpSp"))]
                    if glyphs:
                        for g2 in glyphs:
                            g2["badge"] = badge
                        taken.append(el)
                        icons.extend(glyphs)
                        continue
            taken.append(el)
            icons.append(rec)
    return shapes, texts, slots, icons


def is_ancestor(anc, el):
    p = el.getparent()
    while p is not None:
        if p is anc:
            return True
        p = p.getparent()
    return False


def first_size(tb):
    for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
        for r in tb.iter(q(tag)):
            if r.get("sz"):
                return int(r.get("sz")) / 100
    return None


# ----------------------------------------------------------------------------- inspect

def cmd_inspect(a):
    prs = Presentation(a.library)
    s = lib_slide(prs, a.lib_slide)
    tops, skipped = collect(s)
    shapes, texts, slots, icons = index(tops)
    title = ""
    for sh in s.shapes:
        if sh.is_placeholder and sh.placeholder_format.type == 1:
            title = sh.text_frame.text
    print(f"Library slide {a.lib_slide}: {title!r}   (library canvas 10 x 5.625 in)")
    print("\nTop-level shapes (drop with --drop T#):")
    for top in tops:
        el = top["el"]
        r = top["rect"] if top["slot"] else rect_of(el)
        kind = "IMAGE SLOT" if top["slot"] else el.tag.split("}")[1]
        n_txt = sum(1 for t in texts if t["T"] == top["T"])
        print(f"  T{top['T']:<3} {kind:<12} {name_of(el)[:28]:<28} {fmt_rect(r)}  texts={n_txt}")
    print("\nTexts (set with --texts {\"t#\": \"...\"}):")
    for t in texts:
        sz = f"{t['sz']:.0f}pt" if t["sz"] else "  ?"
        txt = t["text"].replace("\n", " | ")
        print(f"  t{t['t']:<3} T{t['T']:<3} {sz:>5} {fmt_rect(t['abs'])}  {txt[:70]}")
    print("\nIcon candidates (swap with --icons {\"s#\": \"fa-name\"}, remove with --drop s#):")
    for ic in icons:
        print(f"  s{ic['s']:<4} T{ic['T']:<3} {fmt_rect(ic['abs'])}")
    if slots:
        print("\nImage slots (fill in order with --images a.png,b.png; unfilled slots are dropped):")
        for sl in slots:
            print(f"  i{sl['i']:<3} T{sl['top']['T']:<3} {fmt_rect(sl['abs'])} shape={sl['top']['geom']}")
    for sk in skipped:
        print(f"\nSkipped: {sk}")


def fmt_rect(r):
    if not r:
        return "(no xfrm)".ljust(30)
    return f"x{r[0]/EMU_IN:5.2f} y{r[1]/EMU_IN:5.2f} w{r[2]/EMU_IN:5.2f} h{r[3]/EMU_IN:5.2f}"


# ----------------------------------------------------------------------------- edits

def parse_json_arg(val):
    if not val:
        return None
    if os.path.exists(val):
        with open(val, encoding="utf8") as f:
            return json.load(f)
    return json.loads(val)


def set_text(tb, value):
    paras = tb.findall("a:p", NS)
    tpl = next((p for p in paras if p.find("a:r", NS) is not None), paras[0] if paras else None)
    if tpl is None:
        return
    run_tpl = tpl.find("a:r", NS)
    for p in paras:
        tb.remove(p)
    for line in str(value).split("\n"):
        np_ = copy.deepcopy(tpl)
        for ch in list(np_):
            if ch.tag in (q("a:r"), q("a:br"), q("a:fld")):
                np_.remove(ch)
        end = np_.find("a:endParaRPr", NS)
        segs = []
        m = re.match(r"^\*\*(.+?)\*\*(.*)$", line)
        if m:
            segs = [(m.group(1), True), (m.group(2), None)]
        elif line:
            segs = [(line, None)]
        for txt, bold in segs:
            if not txt or run_tpl is None:
                continue
            r = copy.deepcopy(run_tpl)
            r.find("a:t", NS).text = txt
            rpr = r.find("a:rPr", NS)
            if bold is not None:
                if rpr is None:
                    rpr = etree.SubElement(r, q("a:rPr"))
                    r.insert(0, rpr)
                rpr.set("b", "1")
            elif rpr is not None and m:
                rpr.set("b", "0")
            if end is not None:
                end.addprevious(r)
            else:
                np_.append(r)
        tb.append(np_)


def fa_lookup():
    m = {}
    with open(FA_MAP, encoding="utf8") as f:
        for line in f:
            code, names = line.rstrip("\n").split("\t")
            for n in names.split(", "):
                m[n.strip()] = chr(int(code, 16))
    return m


def resolve_color(el, palette, mapping):
    """Best-effort hex for an icon: first solid fill found in it (after palette mapping)."""
    for sf in el.iter(q("a:solidFill")):
        c = sf[0] if len(sf) else None
        if c is None:
            continue
        if c.tag == q("a:srgbClr"):
            return c.get("val")
        if c.tag == q("a:schemeClr"):
            v = c.get("val")
            if v in ACCENTS:
                if palette == "quantzig":
                    return QZ_MAP[v]
                if palette == "custom" and v in mapping:
                    return mapping[v]
                return LIB_ACCENTS[v]
            if v in ("bg1", "lt1"):
                return "FFFFFF"
    return "0B0E5F" if palette == "quantzig" else "404040"


def render_glyph_png(ch, hexcol, px=512):
    from PIL import Image, ImageDraw, ImageFont
    font = ImageFont.truetype(FA_TTF, int(px * 0.86))
    im = Image.new("RGBA", (px, px), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    bb = d.textbbox((0, 0), ch, font=font)
    w, h = bb[2] - bb[0], bb[3] - bb[1]
    d.text(((px - w) / 2 - bb[0], (px - h) / 2 - bb[1]), ch, font=font, fill="#" + hexcol)
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


def make_pic(rid, name, x, y, cx, cy, geom="rect", src=None):
    src_xml = ""
    if src:
        src_xml = "<a:srcRect " + " ".join(f'{k}="{v}"' for k, v in src.items()) + "/>"
    xml = (f'<p:pic xmlns:p="{P}" xmlns:a="{A}" xmlns:r="{R}"><p:nvPicPr><p:cNvPr id="0" name="{name}"/>'
           f'<p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr><p:nvPr/></p:nvPicPr>'
           f'<p:blipFill><a:blip r:embed="{rid}"/>{src_xml}<a:stretch><a:fillRect/></a:stretch></p:blipFill>'
           f'<p:spPr><a:xfrm><a:off x="{int(x)}" y="{int(y)}"/><a:ext cx="{int(cx)}" cy="{int(cy)}"/></a:xfrm>'
           f'<a:prstGeom prst="{geom}"><a:avLst/></a:prstGeom></p:spPr></p:pic>')
    return etree.fromstring(xml)


def make_glyph_box(ch, x, y, cx, cy, color_el, size_pt):
    xml = (f'<p:sp xmlns:p="{P}" xmlns:a="{A}"><p:nvSpPr><p:cNvPr id="0" name="Icon"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
           f'<p:spPr><a:xfrm><a:off x="{int(x)}" y="{int(y)}"/><a:ext cx="{int(cx)}" cy="{int(cy)}"/></a:xfrm>'
           f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
           f'<p:txBody><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" anchor="ctr"><a:noAutofit/></a:bodyPr><a:lstStyle/>'
           f'<a:p><a:pPr algn="ctr"/><a:r><a:rPr lang="en-US" sz="{int(size_pt*100)}" dirty="0"><a:solidFill/>'
           f'<a:latin typeface="FontAwesome"/><a:sym typeface="FontAwesome"/></a:rPr><a:t></a:t></a:r></a:p></p:txBody></p:sp>')
    sp = etree.fromstring(xml)
    sp.find(".//a:t", NS).text = ch
    sf = sp.find(".//a:rPr/a:solidFill", NS)
    sf.append(copy.deepcopy(color_el))
    return sp


def local_rect(el):
    """rect in the coordinate space of el's parent (what we need to put a replacement at)."""
    return rect_of(el)


# ----------------------------------------------------------------------------- colours / fonts

def light_or_dark(clr):
    mods = {c.tag.split("}")[1]: int(c.get("val", "0")) for c in clr}
    if "tint" in mods or mods.get("lumOff", 0) >= 30000:
        return "light"
    if "shade" in mods or (mods.get("lumMod", 100000) <= 80000 and "lumOff" not in mods):
        return "dark"
    return "base"


def alpha_of(clr):
    a = clr.find("a:alpha", NS)
    return a.get("val") if a is not None else None


def qz_gradient(alpha=None):
    al = f'<a:alpha val="{alpha}"/>' if alpha else ""
    return etree.fromstring(
        f'<a:gradFill xmlns:a="{A}" rotWithShape="1"><a:gsLst>'
        f'<a:gs pos="0"><a:srgbClr val="{QZ_WINE}">{al}</a:srgbClr></a:gs>'
        f'<a:gs pos="100000"><a:srgbClr val="{QZ_NAVY}">{al}</a:srgbClr></a:gs>'
        f'</a:gsLst><a:lin ang="10800000" scaled="0"/></a:gradFill>')


def scheme_to_srgb(sc, hexv):
    new = etree.Element(q("a:srgbClr"))
    new.set("val", hexv)
    for ch in sc:
        new.append(copy.deepcopy(ch))
    sc.getparent().replace(sc, new)


def apply_palette(root, palette, mapping):
    if palette == "keep":
        return
    if palette == "quantzig":
        # 0) hard-coded library colours (yellow, orange, blue...) -> brand; greys, white, black kept
        brand = {"0B0E5F", "812B47", "7F2A4C", "2C165E", "592258", "242424", "000000", "FFFFFF",
                 "F8F4F6", "EDD3D4", "D2D0E1"}
        for c in list(root.iter(q("a:srgbClr"))):
            v = (c.get("val") or "").upper()
            if len(v) != 6 or v in brand:
                continue
            r, g, b = int(v[0:2], 16), int(v[2:4], 16), int(v[4:6], 16)
            if max(r, g, b) - min(r, g, b) < 18:
                continue  # grey scale
            lum = (r + g + b) / 3
            c.set("val", "F8F4F6" if lum > 235 else "D2D0E1" if lum > 190 else "592258" if lum > 130 else QZ_NAVY)
        # 1) shape fills -> brand gradient / 4% navy tint / navy shade
        for sppr in list(root.iter(q("p:spPr"))):
            sf = sppr.find("a:solidFill", NS)
            if sf is None or not len(sf):
                continue
            clr = sf[0]
            if clr.tag != q("a:schemeClr") or clr.get("val") not in ACCENTS:
                continue
            kind = light_or_dark(clr)
            if kind == "light":
                mods = {ch.tag.split("}")[1]: int(ch.get("val", "0")) for ch in clr}
                off = mods.get("lumOff", 0) or (100000 - mods.get("tint", 100000))
                al = alpha_of(clr)
                al_xml = f'<a:alpha val="{al}"/>' if al else ""
                tint = "F8F4F6" if off >= 60000 else "D2D0E1"   # keep light/lighter shading steps
                new = etree.fromstring(f'<a:solidFill xmlns:a="{A}"><a:srgbClr val="{tint}">{al_xml}</a:srgbClr></a:solidFill>')
            elif kind == "dark":
                new = etree.fromstring(f'<a:solidFill xmlns:a="{A}"><a:srgbClr val="{QZ_NAVY}"/></a:solidFill>')
            else:
                new = qz_gradient(alpha_of(clr))
            sppr.replace(sf, new)
        mapping = QZ_MAP
        # library greys on text -> Quantzig body text colour
        for rpr in list(root.iter(q("a:rPr"))) + list(root.iter(q("a:defRPr"))) + list(root.iter(q("a:endParaRPr"))):
            sc_ = rpr.find("a:solidFill/a:schemeClr", NS)
            if sc_ is not None and sc_.get("val") in ("tx1", "dk1", "tx2", "dk2"):
                scheme_to_srgb(sc_, "242424")
                for ch in list(rpr.find("a:solidFill/a:srgbClr", NS)):
                    if ch.tag in (q("a:lumMod"), q("a:lumOff"), q("a:tint"), q("a:shade")):
                        ch.getparent().remove(ch)
                continue
            c = rpr.find("a:solidFill/a:srgbClr", NS)
            if c is not None:
                v = c.get("val", "").upper()
                if len(v) == 6 and v[0:2] == v[2:4] == v[4:6] and 0x30 <= int(v[0:2], 16) <= 0xB0:
                    c.set("val", "242424")
    for sc in list(root.iter(q("a:schemeClr"))):
        v = sc.get("val")
        if v in mapping:
            scheme_to_srgb(sc, mapping[v])


def apply_fonts(root, font, scale, min_font):
    for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
        for r in root.iter(q(tag)):
            lat = r.find("a:latin", NS)
            is_icon = lat is not None and "awesome" in (lat.get("typeface") or "").lower()
            if font and not is_icon:
                for sub in ("a:latin", "a:ea", "a:cs"):
                    f = r.find(sub, NS)
                    if f is not None and not (f.get("typeface") or "").startswith("+"):
                        f.set("typeface", font)
            if r.get("sz"):
                sz = int(round(int(r.get("sz")) * scale / 100.0)) * 100   # whole points only
                if min_font and not is_icon and sz < min_font * 100:
                    sz = int(min_font * 100)
                r.set("sz", str(max(sz, 100)))
    if font:  # bullet and symbol faces too (Arial bullets etc.), icon/symbol fonts kept
        for tag in ("a:buFont", "a:sym", "a:latin", "a:ea", "a:cs"):
            for f in root.iter(q(tag)):
                tf = (f.get("typeface") or "")
                low = tf.lower()
                if tf.startswith("+") or "awesome" in low or "wingding" in low or low == "symbol":
                    continue
                f.set("typeface", font)
    for ln in root.iter(q("a:ln")):
        if ln.get("w"):
            ln.set("w", str(int(int(ln.get("w")) * scale)))


# ----------------------------------------------------------------------------- build group

def build_visual(a, target_slide, tgt_prs, box_emu):
    """Returns the finished p:grpSp element (already rel-fixed for target_slide) and a warnings list."""
    lib = Presentation(a.library)
    src = lib_slide(lib, a.lib_slide)
    tops, skipped = collect(src)
    warnings = list(skipped)

    # work on deep copies so indices line up with inspect; images re-linked to the target now
    for top in tops:
        top["el"] = copy.deepcopy(top["el"])
        if not top["slot"]:
            relink_images(top["el"], top["part"], target_slide.part, warnings)
    shapes, texts, slots, icons = index(tops)
    by_s = {f"s{r['s']}": r for r in shapes}
    by_t = {f"t{t['t']}": t for t in texts}

    # --- texts
    tx = parse_json_arg(a.texts)
    if isinstance(tx, list):
        tx = {f"t{i}": v for i, v in enumerate(tx)}
    for k, v in (tx or {}).items():
        if v is None:
            continue
        if k not in by_t:
            warnings.append(f"--texts: no {k} on library slide {a.lib_slide}")
            continue
        set_text(by_t[k]["tb"], v)

    # --- icons (collect replacements before drops)
    icon_req = parse_json_arg(a.icons) or {}
    fa = fa_lookup() if icon_req else {}
    pending_images = []  # (placeholder element, blob)
    for k, name in icon_req.items():
        rec = by_s.get(k)
        if rec is None:
            warnings.append(f"--icons: no {k}")
            continue
        name = name.replace("fa-", "")
        if name not in fa:
            warnings.append(f"--icons: unknown FontAwesome 4 icon '{name}' (see references/fontawesome4-icons.tsv)")
            continue
        el = rec["el"]
        r = local_rect(el)
        if r is None:
            warnings.append(f"--icons: {k} has no position")
            continue
        x, y, cx, cy = r
        abs_h = rec["abs"][3] if rec["abs"] else cy
        if rec.get("badge") and rect_of(rec["badge"]["el"]):
            bx_, by_, bcx, bcy = rect_of(rec["badge"]["el"])       # sibling: same coordinate space
            side = 0.55 * min(bcx, bcy)
            x, y = bx_ + (bcx - side) / 2, by_ + (bcy - side) / 2
            abs_h = rec["badge"]["abs"][3] * 0.55 if rec["badge"]["abs"] else abs_h
        else:
            side = max(cx, cy)
            abs_h = abs_h * side / cy if cy else abs_h
            x, y = x + (cx - side) / 2, y + (cy - side) / 2
        hexcol = a.icon_color or resolve_color(el, a.palette, custom_map(a))
        if a.icon_mode == "glyph":
            color_el = etree.fromstring(f'<a:srgbClr xmlns:a="{A}" val="{hexcol}"/>')
            # glyph size: slide-space height of the icon in pt, times group scale applied later via apply_fonts
            size_pt = abs_h / 12700 * 0.8
            new = make_glyph_box(fa[name], x, y, side, side, color_el, size_pt)
        else:
            new = make_pic("rIdPENDING", f"Icon {name}", x, y, side, side)
            pending_images.append((new, render_glyph_png(fa[name], hexcol)))
        parent = el.getparent()
        if parent is not None:
            parent.replace(el, new)
        else:
            tops[rec["T"]]["el"] = new

    # --- drops
    drops = [d.strip() for d in (a.drop or "").split(",") if d.strip()]
    for d in drops:
        if d.startswith("T"):
            i = int(d[1:])
            if i < len(tops):
                tops[i]["dropped"] = True
        elif d.startswith("s") and d in by_s:
            rec = by_s[d]
            if rec["depth"] == 0:
                tops[rec["T"]]["dropped"] = True
            elif rec["el"].getparent() is not None:
                rec["el"].getparent().remove(rec["el"])
        else:
            warnings.append(f"--drop: unknown label {d}")

    # --- image slots
    imgs = [p for p in (a.images or "").split(",") if p.strip()]
    for sl in slots:
        top = sl["top"]
        if top.get("dropped"):
            continue
        if sl["i"] < len(imgs):
            path = imgs[sl["i"]].strip()
            with open(path, "rb") as f:
                blob = f.read()
            from PIL import Image
            iw, ih = Image.open(io.BytesIO(blob)).size
            x, y, cx, cy = top["rect"]
            src_rect = cover_crop(iw, ih, cx, cy)
            pic = make_pic("rIdPENDING", f"Image {sl['i']}", x, y, cx, cy, top["geom"], src_rect)
            pending_images.append((pic, blob))
            top["el"] = pic
        else:
            top["dropped"] = True

    kept = [t["el"] for t in tops if not t.get("dropped")]
    if not kept:
        sys.exit("nothing left to place")

    # --- bbox of kept tops (library slide space)
    rects = [rect_of(e) for e in kept]
    rects = [r for r in rects if r]
    x0 = min(r[0] for r in rects); y0 = min(r[1] for r in rects)
    x1 = max(r[0] + r[2] for r in rects); y1 = max(r[1] + r[3] for r in rects)
    cw, ch = x1 - x0, y1 - y0
    bx, by, bw, bh = box_emu
    s = min(bw / cw, bh / ch)
    gw, gh = cw * s, ch * s
    gx, gy = bx + (bw - gw) / 2, by + (bh - gh) / 2

    grp = etree.fromstring(
        f'<p:grpSp xmlns:p="{P}" xmlns:a="{A}"><p:nvGrpSpPr><p:cNvPr id="0" name="Visual L{a.lib_slide}"/>'
        f'<p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr><a:xfrm>'
        f'<a:off x="{int(gx)}" y="{int(gy)}"/><a:ext cx="{int(gw)}" cy="{int(gh)}"/>'
        f'<a:chOff x="{x0}" y="{y0}"/><a:chExt cx="{cw}" cy="{ch}"/></a:xfrm></p:grpSpPr></p:grpSp>')
    for e in kept:
        grp.append(e)

    # --- colours and fonts
    apply_palette(grp, a.palette, custom_map(a))
    apply_fonts(grp, a.font, s, a.min_font)

    # --- relationships: images (library) and pending (new)
    for pic, blob in pending_images:
        _, new_rid = target_slide.part.get_or_add_image_part(io.BytesIO(blob))
        b = pic.find(".//a:blip", NS)
        if b is not None:
            b.set(q("r:embed"), new_rid)
    for el in list(grp.iter()):
        if el.tag in (q("a:hlinkClick"), q("a:hlinkHover")):
            el.getparent().remove(el)
            continue
        for att in list(el.attrib):
            if att.startswith(f"{{{R}}}") and el.tag != q("a:blip"):
                del el.attrib[att]
                warnings.append(f"dropped unsupported relationship on <{el.tag.split('}')[1]}>")

    # --- unique ids
    tree = target_slide.shapes._spTree
    used = [int(c.get("id")) for c in tree.iter(q("p:cNvPr")) if c.get("id", "").isdigit()]
    nxt = max(used + [1]) + 1
    for c in grp.iter(q("p:cNvPr")):
        c.set("id", str(nxt)); nxt += 1

    # --- leftover placeholder copy warning
    left = []
    for t in texts:
        if t["tb"].getroottree().getroot() is not None and is_ancestor(grp, t["tb"]):
            tt = text_of(t["tb"])
            if PLACEHOLDER_TEXT.search(tt):
                left.append(f"t{t['t']}")
    if left:
        warnings.append("library placeholder copy still present in: " + ", ".join(left) +
                        "  -> give them real text via --texts or remove them via --drop")
    return grp, warnings, s


def relink_images(el, src_part, target_part, warnings=None):
    """Copy every image an element references from src_part into target_part and repoint r:embed."""
    for blip in list(el.iter(q("a:blip"))):
        rid = blip.get(q("r:embed"))
        if not rid or rid == "rIdPENDING":
            continue
        try:
            blob = src_part.related_part(rid).blob
            _, new_rid = target_part.get_or_add_image_part(io.BytesIO(blob))
            blip.set(q("r:embed"), new_rid)
        except Exception as ex:  # noqa
            if warnings is not None:
                warnings.append(f"image {rid} could not be copied: {ex}")


def custom_map(a):
    m = {}
    for kv in (a.colors or "").split(","):
        if "=" in kv:
            k, v = kv.split("=", 1)
            m[k.strip()] = v.strip().lstrip("#").upper()
    return m


def cover_crop(iw, ih, cx, cy):
    ia, ba = iw / ih, cx / cy
    if abs(ia - ba) < 1e-3:
        return None
    if ia > ba:
        c = int((1 - ba / ia) / 2 * 100000)
        return {"l": c, "r": c}
    c = int((1 - ia / ba) / 2 * 100000)
    return {"t": c, "b": c}


def parse_box(a, prs):
    W, H = prs.slide_width, prs.slide_height
    if a.box:
        x, y, w, h = [float(v) * EMU_IN for v in a.box.split(",")]
        return (x, y, w, h)
    top = 1.6 * EMU_IN if W > 11 * EMU_IN else 0.95 * EMU_IN  # clears the Quantzig header ribbon
    return (0.5 * EMU_IN, top, W - 1.0 * EMU_IN, H - top - 0.6 * EMU_IN)


# ----------------------------------------------------------------------------- place / render

def cmd_place(a):
    prs = Presentation(a.target)
    if a.new_slide is not None:
        slide = prs.slides.add_slide(prs.slide_layouts[a.new_slide])
    else:
        if not a.slide or a.slide > len(prs.slides):
            sys.exit("give --slide N (1-based, existing) or --new-slide LAYOUT_INDEX")
        slide = prs.slides[a.slide - 1]
    if a.title:
        ph = next((p for p in slide.placeholders
                   if p.placeholder_format.type in (1, 3) or p.placeholder_format.idx in (0, 14)), None)
        if ph is not None:
            ph.text_frame.text = a.title
    grp, warnings, s = build_visual(a, slide, prs, parse_box(a, prs))
    tree = slide.shapes._spTree
    ext = tree.find("p:extLst", NS)
    if ext is not None:
        ext.addprevious(grp)
    else:
        tree.append(grp)
    prs.save(a.out)
    idx = list(prs.slides).index(slide) + 1
    print(f"placed library slide {a.lib_slide} on slide {idx} of {a.out} (scale x{s:.2f})")
    for w in warnings:
        print("WARNING:", w)


def soffice(args, cwd):
    wrapper = "/mnt/skills/public/pptx/scripts/office/soffice.py"
    cmd = ([sys.executable, wrapper] if os.path.exists(wrapper) else ["soffice"]) + args
    subprocess.run(cmd, cwd=cwd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=240)


def ensure_fa_font():
    d = os.path.expanduser("~/.fonts")
    dst = os.path.join(d, "fontawesome-webfont.ttf")
    if not os.path.exists(dst):
        os.makedirs(d, exist_ok=True)
        shutil.copy(FA_TTF, dst)
        subprocess.run(["fc-cache", "-f", d], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def render_pptx_to_png(pptx_path, png_path, dpi, crop=True, pad=12):
    from PIL import Image, ImageChops
    ensure_fa_font()
    work = os.path.dirname(pptx_path)
    soffice(["--headless", "--convert-to", "pdf", os.path.basename(pptx_path)], work)
    pdf = os.path.splitext(pptx_path)[0] + ".pdf"
    base = os.path.join(work, "pg")
    subprocess.run(["pdftoppm", "-png", "-r", str(dpi), "-f", "1", "-l", "1", pdf, base], check=True)
    out = [f for f in os.listdir(work) if f.startswith("pg") and f.endswith(".png")][0]
    im = Image.open(os.path.join(work, out)).convert("RGB")
    if crop:
        bg = Image.new("RGB", im.size, (255, 255, 255))
        bb = ImageChops.difference(im, bg).getbbox()
        if bb:
            im = im.crop((max(bb[0] - pad, 0), max(bb[1] - pad, 0),
                          min(bb[2] + pad, im.width), min(bb[3] + pad, im.height)))
    im.save(png_path)
    return im


def blank_deck(width_in, height_in):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(int(width_in * EMU_IN)), Emu(int(height_in * EMU_IN))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    return prs, slide


def cmd_render(a):
    tmp = tempfile.mkdtemp()
    prs, slide = blank_deck(a.width, a.width * 0.5625)
    W, H = prs.slide_width, prs.slide_height
    box = (0.2 * EMU_IN, 0.2 * EMU_IN, W - 0.4 * EMU_IN, H - 0.4 * EMU_IN)
    grp, warnings, s = build_visual(a, slide, prs, box)
    slide.shapes._spTree.append(grp)
    p = os.path.join(tmp, "v.pptx")
    prs.save(p)
    render_pptx_to_png(p, a.out, a.dpi)
    print(f"rendered library slide {a.lib_slide} -> {a.out}")
    for w in warnings:
        print("WARNING:", w)
    if a.keep_pptx:
        shutil.copy(p, a.keep_pptx)


def cmd_preview(a):
    """Render the untouched library slide visual and draw t#/s#/i# labels over it."""
    from PIL import ImageDraw, ImageFont
    tmp = tempfile.mkdtemp()
    lib = Presentation(a.library)
    src = lib_slide(lib, a.lib_slide)
    tops, _ = collect(src)
    shapes, texts, slots, icons = index(tops)
    prs, slide = blank_deck(10, 5.625)
    # place at identity scale, same coordinates, keep image slots as grey frames
    grp = etree.fromstring(
        f'<p:grpSp xmlns:p="{P}" xmlns:a="{A}"><p:nvGrpSpPr><p:cNvPr id="2" name="preview"/>'
        f'<p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/></p:grpSp>')
    for top in tops:
        if top["slot"]:
            x, y, cx, cy = top["rect"]
            grp.append(etree.fromstring(
                f'<p:sp xmlns:p="{P}" xmlns:a="{A}"><p:nvSpPr><p:cNvPr id="3" name="slot"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
                f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
                f'<a:prstGeom prst="{top["geom"]}"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="E4E4E4"/></a:solidFill></p:spPr></p:sp>'))
        else:
            e = copy.deepcopy(top["el"])
            relink_images(e, top["part"], slide.part)
            grp.append(e)
    for el in list(grp.iter()):
        for att in list(el.attrib):
            if att.startswith(f"{{{R}}}") and el.tag != q("a:blip"):
                del el.attrib[att]
    n = 10
    for c in grp.iter(q("p:cNvPr")):
        c.set("id", str(n)); n += 1
    slide.shapes._spTree.append(grp)
    p = os.path.join(tmp, "v.pptx")
    prs.save(p)
    dpi = 150
    im = render_pptx_to_png(p, a.out, dpi, crop=False).convert("RGB")
    d = ImageDraw.Draw(im)
    try:
        f = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 15)
    except Exception:
        f = ImageFont.load_default()
    k = im.width / (10 * EMU_IN)

    def tag(label, ab, color):
        if not ab:
            return
        x, y = ab[0] * k, ab[1] * k
        tw = d.textlength(label, font=f) + 6
        d.rectangle([x, y, x + tw, y + 18], fill=color)
        d.text((x + 3, y + 1), label, fill="white", font=f)

    for ic in icons:
        tag(f"s{ic['s']}", ic["abs"], (200, 0, 120))
    for sl in slots:
        tag(f"i{sl['i']}", sl["abs"], (0, 120, 0))
    for t in texts:
        tag(f"t{t['t']}", t["abs"], (220, 30, 30))
    im.save(a.out)
    print(f"preview -> {a.out}  (red t# = text, magenta s# = icon, green i# = image slot)")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    def common(p, edit=True):
        p.add_argument("lib_slide", type=int, help="1-based slide number in the visual library")
        p.add_argument("--library", default=LIB_DEFAULT)
        if edit:
            p.add_argument("--texts"); p.add_argument("--drop"); p.add_argument("--icons")
            p.add_argument("--icon-mode", choices=["png", "glyph"], default="png")
            p.add_argument("--icon-color", help="hex for swapped icons (default: colour of the icon replaced)")
            p.add_argument("--images")
            p.add_argument("--palette", choices=["keep", "quantzig", "custom"], default="keep")
            p.add_argument("--colors"); p.add_argument("--font")
            p.add_argument("--min-font", type=float, default=0)

    p = sub.add_parser("inspect"); common(p, edit=False); p.set_defaults(fn=cmd_inspect)
    p = sub.add_parser("preview"); common(p, edit=False); p.add_argument("-o", "--out", required=True); p.set_defaults(fn=cmd_preview)
    p = sub.add_parser("place"); common(p)
    p.add_argument("--target", required=True); p.add_argument("--slide", type=int)
    p.add_argument("--new-slide", type=int, help="0-based layout index to add a new slide from")
    p.add_argument("--title"); p.add_argument("--box"); p.add_argument("-o", "--out", required=True)
    p.set_defaults(fn=cmd_place)
    p = sub.add_parser("render"); common(p)
    p.add_argument("-o", "--out", required=True); p.add_argument("--dpi", type=int, default=220)
    p.add_argument("--width", type=float, default=10.0, help="canvas width in inches (16:9)")
    p.add_argument("--keep-pptx"); p.set_defaults(fn=cmd_render)
    a = ap.parse_args()
    if getattr(a, "palette", None) == "quantzig":
        if not a.font:
            a.font = "Grandview"   # Quantzig decks: Grandview is the only text font
        if not a.min_font:
            a.min_font = 10        # Quantzig decks: whole-point sizes from 10pt
    a.fn(a)


if __name__ == "__main__":
    main()
